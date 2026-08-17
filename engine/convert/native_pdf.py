from pathlib import Path

import fitz
from docx import Document
from docx.shared import Inches

from ..ocr import OcrPageResult, ocr_page

_MAX_IMAGE_INCHES = 6.2


def html_from_pdf(pdf_path: Path, html_path: Path) -> int:
    import base64

    doc = fitz.open(str(pdf_path))
    parts: list[str] = ["<!DOCTYPE html><html><head><meta charset='utf-8'></head><body>"]
    images_kept = 0
    try:
        for page_idx, page in enumerate(doc):
            text = page.get_text("text").strip()
            if text:
                parts.append(f"<p>{_html_escape(text)}</p>")
            for i, info in enumerate(page.get_image_info()):
                xref = info.get("xref")
                bbox = info.get("bbox")
                if not bbox:
                    continue
                data = None
                if xref:
                    try:
                        data = doc.extract_image(xref).get("image")
                    except Exception:
                        data = None
                if data is None:
                    pix = page.get_pixmap(clip=fitz.Rect(*bbox), alpha=False)
                    data = pix.tobytes("png")
                    mime = "image/png"
                else:
                    mime = "image/png"
                b64 = base64.b64encode(data).decode("ascii")
                parts.append(f"<img src='data:{mime};base64,{b64}' style='max-width:100%'/><br/>")
                images_kept += 1
    finally:
        doc.close()
    parts.append("</body></html>")
    html_path.write_text("\n".join(parts), encoding="utf-8")
    return images_kept


def _html_escape(text: str) -> str:
    return (
        text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
    )


def _extract_page_images(
    doc: fitz.Document, page_idx: int, workdir: Path, prefix: str, max_images: int = 50
) -> list[tuple[Path, float]]:
    items: list[tuple[Path, float]] = []
    page = doc[page_idx]
    for i, info in enumerate(page.get_image_info()[:max_images]):
        w_pt = info.get("width") or 0
        h_pt = info.get("height") or 0
        bbox = info.get("bbox")
        if w_pt <= 0 or h_pt <= 0 or not bbox:
            continue
        png = workdir / f"{prefix}_p{page_idx}_i{i}.png"
        pix = page.get_pixmap(clip=fitz.Rect(*bbox), alpha=False)
        pix.save(str(png))
        items.append((png, min(w_pt, 72.0 * 6.2)))
    return items


def _ocr_image(image_path: Path) -> OcrPageResult:
    return ocr_page(image_path)


def pdf_to_docx(pdf_path: Path, docx_path: Path, workdir: Path) -> tuple[int, list[bool]]:
    doc = fitz.open(str(pdf_path))
    out = Document()
    ocr_warnings: list[bool] = []
    images_added = 0
    try:
        for page_idx, page in enumerate(doc):
            text = page.get_text("text").strip()
            if text:
                out.add_paragraph(text)
            images = _extract_page_images(doc, page_idx, workdir, "d2w")
            for img, w_pt in images:
                try:
                    out.add_picture(str(img), width=Inches(w_pt / 72.0))
                    images_added += 1
                except Exception:
                    continue
                ocr = _ocr_image(img)
                ocr_warnings.append(ocr.warning)
                blobs = [b.text for b in ocr.blocks if b.text.strip()]
                if blobs:
                    out.add_paragraph("[Ảnh] " + "".join(blobs))
    finally:
        doc.close()
    out.save(str(docx_path))
    return images_added, ocr_warnings