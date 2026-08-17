"""Chuyển đổi sang xlsx/csv bằng thư viện native — không dùng LibreOffice.

LibreOffice trên một số máy không load được filter xuất Calc (xlsx/csv), trong khi
Writer (pdf/docx/txt/html) vẫn chạy bình thường. Các cặp nguồn/đích:
  docx -> xlsx, docx -> csv, xlsx -> csv, xlsx -> xlsx
được xử lý bằng python-docx + openpyxl. Ảnh nhúng không giữ được khi xuất sang bảng
tính (bị bỏ qua; pipeline sẽ ghi chú vào brief).
"""

import csv
import shutil
from pathlib import Path

from docx import Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph


def native_spreadsheet_convert(src: Path, out_dir: Path, target_fmt: str) -> Path:
    from .pipeline import ConvertError

    ext = src.suffix.lower().lstrip(".")
    out = out_dir / f"{src.stem}.{target_fmt}"
    try:
        if ext == "xlsx" and target_fmt == "xlsx":
            shutil.copy2(src, out)
        elif ext == "docx" and target_fmt == "xlsx":
            _docx_to_xlsx(src, out)
        elif ext == "docx" and target_fmt == "csv":
            _docx_to_csv(src, out)
        elif ext == "xlsx" and target_fmt == "csv":
            _xlsx_to_csv(src, out)
        else:
            raise ConvertError(
                f"Chưa hỗ trợ chuyển native {ext} -> {target_fmt.upper()} (dùng đường LibreOffice cho cặp này)"
            )
    except ConvertError:
        raise
    except Exception as exc:
        raise ConvertError(
            f"Chuyển đổi {ext} -> {target_fmt.upper()} thất bại ({type(exc).__name__}): {exc}"
        ) from exc
    if not out.exists() or out.stat().st_size == 0:
        raise ConvertError(f"Chuyển đổi {ext} -> {target_fmt.upper()} không tạo ra được file")
    return out


def _iter_docx_blocks(doc):
    """Duyệt body theo đúng thứ tự trong tài liệu: đoạn (p) và bảng (t)."""
    body = doc._body
    for child in doc.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield "p", Paragraph(child, body)
        elif isinstance(child, CT_Tbl):
            yield "t", Table(child, body)


def _docx_to_xlsx(src: Path, out: Path) -> None:
    from openpyxl import Workbook

    doc = Document(str(src))
    wb = Workbook()
    ws_main = wb.active
    ws_main.title = "Nội dung"
    table_idx = 0
    for kind, block in _iter_docx_blocks(doc):
        if kind == "p":
            text = block.text.strip()
            if text:
                ws_main.append([text])
        else:
            table_idx += 1
            ws = wb.create_sheet(title=f"Bảng {table_idx}")
            for row in block.rows:
                ws.append([cell.text.strip() for cell in row.cells])
    wb.save(str(out))


def _docx_to_csv(src: Path, out: Path) -> None:
    doc = Document(str(src))
    rows: list[list[str]] = []
    for kind, block in _iter_docx_blocks(doc):
        if kind == "p":
            text = block.text.strip()
            if text:
                rows.append([text])
        else:
            for row in block.rows:
                rows.append([cell.text.strip() for cell in row.cells])
    _write_csv(out, rows)


def _xlsx_to_csv(src: Path, out: Path) -> None:
    from openpyxl import load_workbook

    wb = load_workbook(str(src))
    ws = wb[wb.sheetnames[0]]
    rows: list[list[str]] = []
    for row in ws.iter_rows(values_only=True):
        rows.append(["" if value is None else str(value) for value in row])
    _write_csv(out, rows)


def _write_csv(out: Path, rows: list[list[str]]) -> None:
    with open(out, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerows(rows)
